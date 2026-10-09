# -*- coding: utf-8 -*-
"""Tracker helpers for the daily job search.

    python job-search/tracker.py merge-csv FILE.csv   # append CSV rows (same 14 columns)
    python job-search/tracker.py seen URL              # exit 0 if URL already seen, 1 if new
    python job-search/tracker.py add ROW.json          # append one row (dict keyed by column name)
    python job-search/tracker.py skip URL REASON...    # judged and not built: never show it again

Every row added is de-duplicated on Posting URL, and its URL is appended to seen_jobs.txt.
Skipped postings go to seen_jobs.txt too, as "URL  # skipped <date>: reason".
"""

import csv
import datetime
import json
import os
import sys

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
TRACKER = os.path.join(HERE, "Job_Tracker.xlsx")
SEEN = os.path.join(HERE, "seen_jobs.txt")

COLUMNS = ["Date Found", "Job Title", "Company", "Location", "Work Arrangement", "Employment Type",
           "Salary (if listed)", "Posting URL", "Why It Matches", "Resume File", "Cover Letter File",
           "Job Description File", "Status", "Date Applied"]


def seen_urls():
    with open(SEEN) as f:
        return {l.split()[0] for l in f if l.strip() and not l.startswith("#")}


def mark_seen(url):
    if url and url not in seen_urls():
        with open(SEEN, "a") as f:
            f.write(url + "\n")


def add_rows(rows):
    wb = openpyxl.load_workbook(TRACKER)
    ws = wb.active
    header = [c.value for c in ws[1]]
    assert header == COLUMNS, f"tracker columns changed: {header}"
    existing = {r[7] for r in ws.iter_rows(min_row=2, values_only=True)}
    added = 0
    for row in rows:
        url = row.get("Posting URL")
        if url in existing:
            continue
        ws.append([row.get(c) for c in COLUMNS])
        existing.add(url)
        mark_seen(url)
        added += 1
    wb.save(TRACKER)
    return added


if __name__ == "__main__":
    cmd = sys.argv[1]
    if cmd == "merge-csv":
        with open(sys.argv[2], newline="") as f:
            print("added", add_rows(list(csv.DictReader(f))))
    elif cmd == "seen":
        sys.exit(0 if sys.argv[2] in seen_urls() else 1)
    elif cmd == "add":
        with open(sys.argv[2]) as f:
            print("added", add_rows([json.load(f)]))
    elif cmd == "skip":
        if sys.argv[2] not in seen_urls():
            with open(SEEN, "a") as f:
                f.write(f"{sys.argv[2]}  # skipped {datetime.date.today()}: {' '.join(sys.argv[3:])}\n")
    else:
        sys.exit(__doc__)
